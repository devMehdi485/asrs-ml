"""Génère la présentation PowerPoint (thème sombre AeroInsight)."""
import os, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "report", "figures")
data = json.load(open(os.path.join(ROOT, "frontend", "public", "data", "asrs.json"), encoding="utf-8"))
w2v = json.load(open(os.path.join(ROOT, "report", "word2vec.json"), encoding="utf-8"))

NAVY = RGBColor(0x0A, 0x0F, 0x1C); CARD = RGBColor(0x12, 0x1A, 0x2E)
BLUE = RGBColor(0x3B, 0x82, 0xF6); WHITE = RGBColor(0xF1, 0xF5, 0xFB)
MUTED = RGBColor(0x9A, 0xA8, 0xC0); GREEN = RGBColor(0x22, 0xC5, 0x5E); LINE = RGBColor(0x24, 0x34, 0x4C)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def slide(bg=NAVY):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    return s


def box(s, x, y, w, h):
    return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame


def settext(tf, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Calibri"
    return p


def header(s, title, sub=None):
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), W, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    settext(box(s, 0.6, 0.35, 12, 1), title, 30, WHITE, True)
    if sub:
        settext(box(s, 0.62, 1.15, 12, 0.6), sub, 15, MUTED)


def bullets(s, items, x=0.8, y=1.7, w=11.7, h=5, size=18, gap=10):
    tf = box(s, x, y, w, h); tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        lvl = 0
        if isinstance(it, tuple): it, lvl = it
        r = p.add_run(); r.text = ("•  " if lvl == 0 else "–  ") + it
        r.font.size = Pt(size - lvl * 2); r.font.color.rgb = WHITE if lvl == 0 else MUTED
        r.font.name = "Calibri"; p.level = lvl


def picture(s, name, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    s.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), **kw)


def kpibar(s, items, y=5.7):
    n = len(items); gap = 0.25; tw = (13.333 - 1.2 - gap * (n - 1)) / n
    for i, (val, lab) in enumerate(items):
        x = 0.6 + i * (tw + gap)
        c = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(tw), Inches(1.2))
        c.fill.solid(); c.fill.fore_color.rgb = CARD; c.line.color.rgb = LINE; c.line.width = Pt(1)
        tf = c.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        settext(tf, val, 26, BLUE, True, PP_ALIGN.CENTER)
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = lab; r.font.size = Pt(11); r.font.color.rgb = MUTED


def table(s, rows, x, y, w, h, header_fill=BLUE, fs=12):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for j in range(nc):
        for i in range(nr):
            cell = gt.cell(i, j); cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (CARD if i % 2 else RGBColor(0x16, 0x22, 0x3A))
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; r = p.add_run(); r.text = str(rows[i][j])
            r.font.size = Pt(fs); r.font.color.rgb = WHITE
            r.font.bold = (i == 0); r.font.name = "Calibri"
    return gt


m = data["meta"]

# 1 TITLE
s = slide()
band = s.shapes.add_shape(1, Inches(0), Inches(2.4), W, Inches(2.7))
band.fill.solid(); band.fill.fore_color.rgb = CARD; band.line.fill.background()
settext(box(s, 0.8, 0.7, 12, 0.5), "AÉRONAUTIQUE · PROJET A3 · NLP · CLUSTERING", 14, BLUE, True)
settext(box(s, 0.8, 2.55, 12, 1.5), "Analyse NLP des rapports d'incidents\nde sécurité aérienne", 40, WHITE, True)
settext(box(s, 0.85, 4.35, 12, 0.6), "NASA ASRS — extraction des causes, classification, signaux faibles", 18, MUTED)
kpibar(s, [(f"{m['n_reports']:,}".replace(",", " "), "rapports analysés"),
           (str(m["n_themes"]), "thèmes détectés"), ("11", "catégories"),
           (f"{m['period'][0][:4]}–{m['period'][1][:4]}", "période")], y=5.7)

# 2 Contexte
s = slide(); header(s, "Contexte & problématique")
bullets(s, [
    "Des milliers de rapports d'incidents soumis volontairement à la NASA (ASRS) chaque année.",
    "Chaque rapport = un récit en texte libre + métadonnées codées par des experts.",
    "Une mine d'informations sur les risques systémiques… largement inexploitée par l'IA.",
    "Problématique : faire parler automatiquement ces récits pour en extraire des causes,",
    ("classifier les incidents et détecter des signaux faibles précurseurs d'accidents.", 1),
])

# 3 Objectifs
s = slide(); header(s, "Objectifs")
bullets(s, [
    "Pipeline NLP complet : prétraitement → vectorisation → clustering.",
    "Faire émerger des thèmes d'incidents cohérents et interprétables (non supervisé).",
    "Classifier le type d'anomalie à partir du texte (supervisé).",
    "Repérer les signaux faibles et analyser l'évolution temporelle.",
    "Livrer un dashboard interactif de qualité professionnelle.",
])

# 4 Données
s = slide(); header(s, "Données — NASA ASRS")
bullets(s, [
    "125 763 rapports réels (2002–2025), accès libre.",
    "Texte : Narratives Reporter 1 & 2 + Synopsis (100 % remplis).",
    "Cibles/métadonnées : Anomalie 99,7 % · Phase de vol 95,9 % · Aéronef 98,9 % · Date 100 %.",
    "Particularité : double ligne d'en-tête gérée automatiquement au chargement.",
], y=1.7, h=3)
kpibar(s, [("125 763", "rapports"), ("126", "variables"), ("100 %", "narratives"),
           ("99,7 %", "anomalie remplie")], y=5.6)

# 5 Méthodologie
s = slide(); header(s, "Méthodologie — pipeline NLP")
bullets(s, [
    "Prétraitement : nettoyage désid. (ZZZ), tokenisation, stopwords métier, lemmatisation.",
    "Vectorisation : TF-IDF · word2vec (gensim) · embeddings de phrases (MiniLM).",
    "Réduction UMAP → clustering K-Means vs HDBSCAN (+ garde-fou anti-effondrement).",
    "Interprétation : c-TF-IDF, labels métier (11 catégories), LDA + cohérence c_v.",
    "Classification supervisée (LogReg, SVM) · Signaux faibles (Isolation Forest) · Temporel.",
])

# 6 Cartographie
s = slide(); header(s, "Cartographie thématique — 81 thèmes, 11 catégories")
picture(s, "categories.png", 0.5, 1.6, w=6.5)
picture(s, "scatter.png", 7.1, 1.5, w=6.0)

# 7 Clustering comparison
s = slide(); header(s, "Cohérence des clusters — K-Means vs HDBSCAN")
table(s, [["Modèle", "Clusters", "Silhouette", "Davies-B.", "Bruit", "Pureté"],
          ["K-Means", "15", "0,382", "0,945", "0 %", "0,133"],
          ["HDBSCAN ✓", "81", "0,590", "0,575", "32 %", "0,165"]],
      x=0.7, y=1.8, w=12, h=2, fs=14)
bullets(s, [
    "HDBSCAN retenu : meilleure silhouette (0,590) et Davies-Bouldin plus faible.",
    "Le bruit (32 %) est réutilisé pour la détection de signaux faibles.",
    "Garde-fou : un HDBSCAN dégénéré (< 6 clusters) est rejeté au profit de K-Means.",
], y=4.2, h=2.6, size=16)

# 8 Wordclouds
s = slide(); header(s, "Nuages de mots par thème (c-TF-IDF)")
picture(s, "wordclouds.png", 1.6, 1.6, w=10)

# 9 word2vec
s = slide(); header(s, "Embeddings de mots (word2vec) — sémantique du domaine")
rows = [["Mot", "Voisins les plus proches (similarité)"]]
for k in ["runway", "engine", "fuel", "weather", "altitude", "gear"]:
    if k in w2v:
        rows.append([k, ", ".join(f"{a} ({b})" for a, b in w2v[k][:5])])
table(s, rows, x=0.7, y=1.8, w=12, h=3.2, fs=13)
settext(box(s, 0.7, 5.4, 12, 0.8),
        "word2vec capture une sémantique fine : runway≈rwy/taxiway, weather≈thunderstorm/storm…", 14, MUTED)

# 10 Classification
s = slide(); header(s, "Classification supervisée du type d'incident")
table(s, [["Modèle", "F1 macro", "F1 pondéré"],
          ["Régression logistique", "0,518", "0,606"],
          ["SVM linéaire", "0,517", "0,615"]], x=0.7, y=1.8, w=8, h=1.8, fs=14)
bullets(s, [
    "TF-IDF + top-8 anomalies les plus fréquentes.",
    "Meilleures classes : ATC Issue (F1 0,76), NMAC (0,69), Equipment Critical (0,67).",
    "Difficultés : classes rares et anomalies multi-étiquettes.",
], y=4.0, h=2.6, size=16)

# 11 Temporel
s = slide(); header(s, "Évolution temporelle des incidents")
picture(s, "temporal_volume.png", 0.5, 1.6, w=6.4)
picture(s, "heatmap.png", 7.0, 1.6, w=6.1)

# 12 Causes + signaux
s = slide(); header(s, "Causes récurrentes & signaux faibles")
table(s, [["Thème", "Rapports", "%", "Tendance"],
          ["Pannes avioniques & systèmes", "22 258", "26,2", "baisse"],
          ["Incursions & erreurs au sol", "8 458", "10,0", "hausse"],
          ["Fumée / odeurs en cabine", "3 393", "4,0", "hausse"],
          ["Conflits de trafic (circuit)", "2 657", "3,1", "hausse"],
          ["Marchandises dangereuses", "1 959", "2,3", "hausse"]],
      x=0.7, y=1.8, w=9.2, h=3, fs=13)
bullets(s, [
    f"{m['n_weak']} rapports atypiques (Isolation Forest).",
    "En hausse : sol, fumées, drones, hazmat.",
    "→ priorités d'un programme sécurité.",
], x=10.2, y=1.9, w=3, h=3, size=13)

# 13 Dashboard
s = slide(); header(s, "Dashboards interactifs")
bullets(s, [
    "Streamlit (Python) — rapide, lit directement les artefacts.",
    "AeroInsight AI (React + Tailwind) — 6 pages, design pro :",
    ("Dashboard · Cluster Explorer (carte + filtres catégories) · Thematic Analysis", 1),
    ("Temporal Map · Weak Signals · Executive Intelligence (rapport auto)", 1),
    "Chaque thème : nom métier, description, nuage de mots, tendance.",
], y=1.7, h=3.2)
kpibar(s, [("2", "dashboards"), ("6", "pages React"), ("81", "thèmes nommés"), ("11", "catégories")], y=5.6)

# 14 Conclusion
s = slide(); header(s, "Conclusion")
bullets(s, [
    "Pipeline NLP complet validé sur 125 763 rapports réels.",
    "81 thèmes interprétables en 11 catégories métier.",
    "Classification opérationnelle (F1 pondéré 0,615) + signaux faibles + temporel.",
    "Deux dashboards transformant des récits bruts en intelligence de sécurité.",
], y=1.7, h=2.6)
settext(box(s, 0.8, 5.0, 12, 0.6), "github.com/devMehdi485/asrs-ml", 16, BLUE, True)

out = os.path.join(ROOT, "report", "presentation.pptx")
prs.save(out)
print("Présentation générée :", out, "|", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
